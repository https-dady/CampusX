from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from pypdf import PdfReader


SECTION_ALIASES = {
    "summary": [
        "summary",
        "profile",
        "objective",
        "professional summary",
    ],
    "education": [
        "education",
        "academic background",
        "academics",
    ],
    "experience": [
        "experience",
        "work experience",
        "professional experience",
        "internship",
        "internships",
    ],
    "skills": [
        "skills",
        "technical skills",
        "technical skill",
        "skills & technologies",
        "technologies",
    ],
    "projects": [
        "projects",
        "project",
    ],
    "certifications": [
        "certifications",
        "certificates",
        "certification",
    ],
    "achievements": [
        "achievements",
        "accomplishments",
    ],
    "contact": [
        "contact",
        "contact information",
    ],
}


def extract_text(
    file_bytes: bytes,
    mimetype: str,
) -> str:
    if mimetype == "application/pdf":
        reader = PdfReader(
            io.BytesIO(file_bytes)
        )

        pages = []

        for page in reader.pages:
            pages.append(
                page.extract_text() or ""
            )

        return "\n".join(
            pages
        ).strip()

    if mimetype == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        document = Document(
            io.BytesIO(file_bytes)
        )

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(
                        cell.text
                        for cell in row.cells
                    )
                )

        return "\n".join(
            parts
        ).strip()

    raise ValueError(
        "Only PDF and DOCX resumes are supported."
    )


def normalize_text(
    value: str,
) -> str:
    return " ".join(
        (value or "").strip().split()
    )


def _clean_lines(
    text: str,
) -> list[str]:
    return [
        normalize_text(line)
        for line in text.splitlines()
        if normalize_text(line)
    ]


def detect_sections(
    text: str,
) -> dict[str, str]:
    lines = _clean_lines(text)

    sections: dict[
        str,
        list[str],
    ] = {}

    current = "other"

    sections[current] = []

    alias_to_section = {
        alias.lower(): section
        for section, aliases in SECTION_ALIASES.items()
        for alias in aliases
    }

    for line in lines:
        key = re.sub(
            r"[^a-z& ]",
            "",
            line.lower(),
        ).strip()

        section = alias_to_section.get(
            key
        )

        if section:
            current = section

            sections.setdefault(
                current,
                [],
            )

            continue

        sections.setdefault(
            current,
            [],
        ).append(line)

    return {
        section: "\n".join(values).strip()
        for section, values in sections.items()
        if values
    }


def _first_match(
    patterns: list[str],
    text: str,
    flags=re.I,
) -> str | None:
    for pattern in patterns:
        match = re.search(
            pattern,
            text,
            flags,
        )

        if match:
            return normalize_text(
                match.group(1)
            )

    return None


def extract_cgpa(
    text: str,
) -> float | None:
    value = _first_match(
        [
            r"(?:cgpa|c\.g\.p\.a)\s*[:\-]?\s*(\d{1,2}(?:\.\d{1,2})?)",
            r"(?:gpa)\s*[:\-]?\s*(\d{1,2}(?:\.\d{1,2})?)",
        ],
        text,
    )

    if value is None:
        return None

    number = float(value)

    return (
        number
        if 0 <= number <= 10
        else None
    )


def extract_academic_year(
    text: str,
) -> int | None:
    patterns = [
        r"\b(1st|first)\s+year\b",
        r"\b(2nd|second)\s+year\b",
        r"\b(3rd|third)\s+year\b",
        r"\b(4th|fourth)\s+year\b",
        r"\b(5th|fifth)\s+year\b",
        r"\b(6th|sixth)\s+year\b",
    ]

    values = [
        1,
        2,
        3,
        4,
        5,
        6,
    ]

    for pattern, value in zip(
        patterns,
        values,
    ):
        if re.search(
            pattern,
            text,
            re.I,
        ):
            return value

    return None


def extract_degree(
    text: str,
) -> str | None:
    patterns = [
        r"\b(B\.\s?Tech|BTech|B\.E\.|BE|B\.Sc|BSc|BCA|BBA|M\.Tech|MTech|MCA|MBA)\b",
    ]

    return _first_match(
        patterns,
        text,
    )


def extract_branch(
    text: str,
) -> str | None:
    branch_patterns = [
        (
            r"artificial intelligence\s*(?:and|&)\s*data science",
            "AI & DS",
        ),
        (
            r"ai\s*(?:and|&)\s*ds",
            "AI & DS",
        ),
        (
            r"computer science(?: and engineering| & engineering| engineering)?",
            "CSE",
        ),
        (
            r"information technology",
            "IT",
        ),
        (
            r"data science",
            "Data Science",
        ),
        (
            r"electronics(?: and| &)\s*communication",
            "ECE",
        ),
        (
            r"business analytics",
            "Business Analytics",
        ),
        (
            r"computer applications",
            "Computer Applications",
        ),
    ]

    for pattern, branch in branch_patterns:
        if re.search(
            pattern,
            text,
            re.I,
        ):
            return branch

    return None


def extract_internship(
    text: str,
) -> bool:
    return bool(
        re.search(
            r"\bintern(ship|ships)?\b",
            text,
            re.I,
        )
    )


def extract_name(
    text: str,
) -> str | None:
    lines = _clean_lines(text)[:8]

    for line in lines:
        if any(
            char.isdigit()
            for char in line
        ):
            continue

        if (
            "@"
            in line
            or "http"
            in line.lower()
        ):
            continue

        if (
            len(line.split())
            in (2, 3, 4)
            and len(line) <= 60
        ):
            return line

    return None


def extract_email(
    text: str,
) -> str | None:
    match = re.search(
        r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b",
        text,
        re.I,
    )

    return (
        match.group(0)
        if match
        else None
    )


def extract_phone(
    text: str,
) -> str | None:
    match = re.search(
        r"(?:\+?\d[\d\s().-]{8,}\d)",
        text,
    )

    return (
        normalize_text(
            match.group(0)
        )
        if match
        else None
    )


def extract_skills(
    text: str,
    known_skills: list[str],
) -> list[str]:
    lower = text.lower()

    found: list[str] = []

    for skill in known_skills:
        normalized = normalize_text(
            str(skill)
        )

        if not normalized:
            continue

        escaped = re.escape(
            normalized.lower()
        )

        if re.search(
            rf"(?<![a-z0-9+#.]){escaped}(?![a-z0-9+#.])",
            lower,
        ):
            found.append(
                normalized
            )

    return sorted(
        set(found),
        key=str.lower,
    )


def build_resume_data(
    text: str,
    known_skills: list[str],
) -> dict[str, Any]:
    sections = detect_sections(
        text
    )

    return {
        "personal": {
            "name": extract_name(text),
            "email": extract_email(text),
            "phone": extract_phone(text),
        },
        "education": {
            "degree": extract_degree(text),
            "branch": extract_branch(text),
            "academicYear": extract_academic_year(
                text
            ),
            "cgpa": extract_cgpa(text),
        },
        "hasInternship": extract_internship(
            text
        ),
        "skills": extract_skills(
            text,
            known_skills,
        ),
        "sections": {
            key: bool(
                sections.get(key)
            )
            for key in SECTION_ALIASES
        },
        "sectionText": sections,
        "textLength": len(text),
        "wordCount": len(
            text.split()
        ),
    }